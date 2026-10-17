"""Hand-label ("gold") set for validating the two LLM judges.

make_sheet: picks answers where the judges disagree plus a few where they agree, shuffles them so the
labeller cannot tell which is which, hides both judges' codings, and writes an Excel sheet with
drop-down cells. score: reads the filled sheet back and reports each judge's agreement with the human.
"""
from __future__ import annotations

import random
import re
from pathlib import Path

import pandas as pd

from . import build, metrics

VERDICTS = ["stay", "switch", "mixed", "n/a"]

GUIDE = """Labelling guide (read only the "conversation" and "answer" columns; fill in the two yellow columns)

top_pick - apply the rules in order:
1. Explicit recommendation ("I'd start with X" / "get X" / "X is the best pick") -> that brand
2. Ranked list -> the rank-1 brand
3. Unranked list -> none
4. A pick with caveats ("X, but it's salty") -> still that brand
5. A pick plus alternatives ("X is best; Y is the budget option") -> the primary brand X
6. Different brands for different conditions ("for protein X, for budget Y") with no overall pick -> conditional
7. Refuses to recommend any brand -> refusal
8. Only a product is named -> its brand
9. Multi-turn: if the new answer does not clearly recommend a different brand -> keep the earlier recommended brand
Store brands (Kirkland, Great Value, Good & Gather, Trader Joe's, Friendly Farms, ...) are always Private label.
No brand mentioned at all -> none.
The answer only reports someone else's ranking ("Food Republic ranked X first") -> use your judgement:
does the assistant present it as its own recommendation?

held_brand_verdict - fill in only when the shopper says "I usually buy <brand>"; leave n/a on all other rows:
stay = advises keeping the brand (another product line of the same brand still counts as stay);
switch = advises a different brand; mixed = gives both without a lean.

notes (optional): one line when unsure. About 50 rows, roughly a minute each.
"""


def _conversation(r: dict) -> str:
    msgs = r["messages"]
    if len(msgs) == 1:
        return msgs[0]["content"]
    parts = []
    for m in msgs:
        if m["role"] == "assistant":
            parts.append(f"ASSISTANT: {m['content']}")
        else:
            parts.append(f"USER: {m['content']}")
    return "\n\n".join(parts)


def _clean(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text or "")


def make_sheet(answers: list[dict], first: dict, second: dict, cfg: dict, out: Path,
               n_disagree: int = 40, n_agree: int = 10, seed: int = 7) -> pd.DataFrame:
    canon = lambda b: build.canonical(b or "none", cfg["brands"])
    by_id = {r["task_id"]: r for r in answers}
    ids = sorted(set(first) & set(second) & set(by_id))

    def differs(t):
        a, b = first[t]["coded"], second[t]["coded"]
        return (canon(a.get("top_pick")) != canon(b.get("top_pick"))
                or a.get("held_brand_verdict", "n/a") != b.get("held_brand_verdict", "n/a"))

    dis = [t for t in ids if differs(t)]
    agree = [t for t in ids if not differs(t)]
    rng = random.Random(seed)
    picked = rng.sample(dis, min(n_disagree, len(dis))) + rng.sample(agree, min(n_agree, len(agree)))
    rng.shuffle(picked)

    rows = []
    for i, t in enumerate(picked, start=1):
        r = by_id[t]
        conv = _conversation(r)
        loyalty = "usually buy" in conv.lower()
        rows.append({"row": i, "module": r["module"], "engine_hidden": r["engine"],
                     "conversation": _clean(conv), "answer": _clean(r["text"]),
                     "top_pick": "", "held_brand_verdict": "" if loyalty else "n/a", "notes": "",
                     "task_id": t, "subset": "disagree" if t in dis else "agree"})
    df = pd.DataFrame(rows)
    _write_xlsx(df, cfg, out)
    return df


def _write_xlsx(df: pd.DataFrame, cfg: dict, out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "label"
    cols = ["row", "module", "conversation", "answer", "top_pick", "held_brand_verdict", "notes", "task_id"]
    widths = [6, 14, 45, 95, 18, 18, 30, 18]
    ws.append(cols)
    for _, r in df.iterrows():
        ws.append([r[c] for c in cols])
    bold, yellow = Font(bold=True), PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths, start=1):
        letter = ws.cell(1, j).column_letter
        ws.column_dimensions[letter].width = w
        ws.cell(1, j).font = bold
    for i in range(2, len(df) + 2):
        ws.row_dimensions[i].height = 300
        for j in range(1, len(cols) + 1):
            ws.cell(i, j).alignment = wrap
        for j in (5, 6):
            ws.cell(i, j).fill = yellow
    ws.column_dimensions["H"].hidden = True
    ws.freeze_panes = "C2"

    lists = wb.create_sheet("lists")
    picks = list(cfg["brands"]) + ["conditional", "none", "refusal"]
    for i, v in enumerate(picks, start=1):
        lists.cell(i, 1, v)
    for i, v in enumerate(VERDICTS, start=1):
        lists.cell(i, 2, v)
    lists.sheet_state = "hidden"
    dv1 = DataValidation(type="list", formula1=f"=lists!$A$1:$A${len(picks)}", allow_blank=True)
    dv2 = DataValidation(type="list", formula1=f"=lists!$B$1:$B${len(VERDICTS)}", allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"E2:E{len(df) + 1}")
    dv2.add(f"F2:F{len(df) + 1}")

    guide = wb.create_sheet("guide", 0)
    for i, line in enumerate(GUIDE.splitlines(), start=1):
        guide.cell(i, 1, line)
    guide.column_dimensions["A"].width = 110
    guide.cell(1, 1).font = bold
    wb.active = 1
    wb.save(out)


def score(sheet: Path, key: pd.DataFrame, judges: dict[str, dict], cfg: dict) -> pd.DataFrame:
    """Compare each judge with the human labels. key: the make_sheet frame (task_id, subset)."""
    canon = lambda b: build.canonical(str(b) if b and str(b) != "nan" else "none", cfg["brands"])
    retailers = {"kroger", "walmart", "target", "costco", "aldi", "whole foods", "trader joe's", "safeway"}

    def canon_human(x):  # typed labels: trim spaces; a retailer name means its store brand
        x = str(x).strip()
        return "Private label" if x.lower() in retailers else canon(x)

    human = pd.read_excel(sheet, sheet_name="label")
    human = human[human.top_pick.notna() & (human.top_pick.astype(str).str.strip() != "")]
    if human.empty:
        raise SystemExit(f"no labels filled in yet in {sheet}")
    human = human.merge(key[["task_id", "subset"]], on="task_id", how="left")
    print(f"\n{len(human)} hand-labelled answers "
          f"({(human.subset == 'disagree').sum()} where the judges disagree, {(human.subset == 'agree').sum()} where they agree)")
    out = []
    for name, coded in judges.items():
        h = human.assign(judge_pick=[canon(coded[t]["coded"].get("top_pick")) if t in coded else None
                                     for t in human.task_id],
                         judge_verdict=[coded[t]["coded"].get("held_brand_verdict", "n/a") if t in coded else None
                                        for t in human.task_id],
                         human_pick=[canon_human(x) for x in human.top_pick])
        res = {"judge": name}
        for subset, part in [("all", h), ("disagree", h[h.subset == "disagree"]), ("agree", h[h.subset == "agree"])]:
            if len(part):
                res[f"top_pick_match_{subset}"] = round((part.judge_pick == part.human_pick).mean(), 2)
        res["top_pick_kappa_all"] = round(metrics.cohen_kappa(list(h.human_pick), list(h.judge_pick)), 2)
        v = h[h.held_brand_verdict.astype(str).isin(["stay", "switch", "mixed"])]
        if len(v):
            res["verdict_match"] = round((v.judge_verdict == v.held_brand_verdict).mean(), 2)
            res["verdict_n"] = len(v)
        out.append(res)
    table = pd.DataFrame(out).set_index("judge")
    # The sample over-represents disagreements; re-weight to the full data set for an accuracy estimate.
    names = list(judges)
    if len(names) == 2:
        a, b = judges[names[0]], judges[names[1]]
        ids = set(a) & set(b)
        dis_rate = sum(canon(a[t]["coded"].get("top_pick")) != canon(b[t]["coded"].get("top_pick"))
                       or a[t]["coded"].get("held_brand_verdict", "n/a") != b[t]["coded"].get("held_brand_verdict", "n/a")
                       for t in ids) / len(ids)
        if {"top_pick_match_agree", "top_pick_match_disagree"} <= set(table.columns):
            table["est_top_pick_accuracy_full_data"] = (
                (1 - dis_rate) * table.top_pick_match_agree + dis_rate * table.top_pick_match_disagree).round(2)
            print(f"(judges disagree on {dis_rate:.0%} of all answers; accuracy re-weighted accordingly)")
    print(table.T.to_string())
    print("\nOn the 'agree' rows both judges give the same answer, so they check that agreement means correct;")
    print("on the 'disagree' rows the higher match rate shows which judge is right more often.")
    return table
